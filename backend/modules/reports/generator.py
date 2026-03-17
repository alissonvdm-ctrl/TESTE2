"""
ReportGenerator – build structured report payloads and render them to
JSON, HTML, PDF, or Excel.

This module is a standalone report builder decoupled from Celery.  The
``generate_report`` Celery task in ``backend/tasks/report_tasks.py`` uses
the higher-level helpers there; this module exposes the core rendering
logic for reuse and testing.
"""
from __future__ import annotations

import json
import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Report payload builder
# ---------------------------------------------------------------------------


def build_report_payload(
    experiment_id: int,
    experiment_name: str,
    experiment_status: str,
    n_max: int,
    k_count: int,
    order_matters: bool,
    description: Optional[str],
    analysis_modules: Dict[str, Dict[str, Any]],
    models: List[Dict[str, Any]],
    config_json: Optional[Any] = None,
) -> Dict[str, Any]:
    """
    Assemble the canonical report payload dict.

    Parameters
    ----------
    experiment_id:
        DB primary key of the experiment.
    experiment_name:
        Human-readable name.
    experiment_status:
        Lifecycle status string (e.g. ``'completed'``).
    n_max, k_count, order_matters:
        Domain configuration.
    description:
        Optional free-text description.
    analysis_modules:
        Dict ``{module_name: {status, duration_seconds, created_at, result}}``.
    models:
        List of model result dicts with keys:
        ``model_name, module, val_score, test_score, trained_at, metrics,
        feature_importance``.
    config_json:
        Arbitrary additional configuration.

    Returns
    -------
    dict
        Structured payload ready for JSON / HTML / Excel rendering.
    """
    now_iso = datetime.now(tz=timezone.utc).isoformat()

    scored = [m for m in models if m.get("val_score") is not None]
    best_model = min(scored, key=lambda m: m["val_score"]) if scored else None

    ensemble_result = analysis_modules.get("ensemble", {}).get("result") or {}

    return {
        "report_metadata": {
            "generated_at": now_iso,
            "experiment_id": experiment_id,
            "experiment_name": experiment_name,
            "experiment_status": experiment_status,
        },
        "experiment_config": {
            "n_max": n_max,
            "k_count": k_count,
            "order_matters": order_matters,
            "description": description,
            "config_json": config_json,
        },
        "summary": {
            "modules_run": list(analysis_modules.keys()),
            "models_trained": [m["model_name"] for m in models],
            "best_model": best_model["model_name"] if best_model else None,
            "best_val_score": best_model["val_score"] if best_model else None,
            "ensemble_prediction": ensemble_result.get("predicted_numbers"),
            "ensemble_confidence": ensemble_result.get("confidence"),
        },
        "analysis": analysis_modules,
        "models": models,
    }


# ---------------------------------------------------------------------------
# Renderers
# ---------------------------------------------------------------------------


def render_json(payload: Dict[str, Any], output_path: Path) -> None:
    """Write the payload as indented JSON."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, indent=2, ensure_ascii=False, default=str)
    logger.info("JSON report written to %s", output_path)


def render_html(payload: Dict[str, Any], output_path: Path) -> None:
    """Render a self-contained HTML report."""
    meta = payload["report_metadata"]
    cfg = payload["experiment_config"]
    summary = payload["summary"]
    analysis = payload["analysis"]
    models = payload["models"]

    def _fmt(v: Any) -> str:
        if v is None:
            return "<em>N/A</em>"
        if isinstance(v, float):
            return f"{v:.4f}"
        if isinstance(v, list):
            return ", ".join(str(x) for x in v)
        return str(v)

    model_rows = "".join(
        f"<tr><td>{m['model_name']}</td>"
        f"<td>{_fmt(m.get('val_score'))}</td>"
        f"<td>{_fmt(m.get('test_score'))}</td>"
        f"<td>{_fmt(m.get('trained_at'))}</td></tr>"
        for m in models
    )
    analysis_rows = "".join(
        f"<tr><td>{name}</td>"
        f"<td>{_fmt(data.get('status'))}</td>"
        f"<td>{_fmt(data.get('duration_seconds'))}</td></tr>"
        for name, data in analysis.items()
    )

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8"/>
  <title>Report – {meta['experiment_name']}</title>
  <style>
    body{{font-family:sans-serif;margin:2rem;color:#222}}
    h1{{color:#1a56db}}h2{{border-bottom:2px solid #e5e7eb;padding-bottom:.3rem}}
    table{{border-collapse:collapse;width:100%;margin-bottom:1.5rem}}
    th,td{{border:1px solid #d1d5db;padding:.5rem .75rem;text-align:left}}
    th{{background:#f3f4f6}}
    .badge{{display:inline-block;padding:.2rem .6rem;border-radius:9999px;
            background:#dbeafe;color:#1e40af;font-size:.85rem}}
  </style>
</head>
<body>
  <h1>Analysis Report – {meta['experiment_name']}</h1>
  <p>Generated: {meta['generated_at']} &nbsp;|&nbsp; Status:
     <span class="badge">{meta.get('experiment_status') or 'unknown'}</span></p>
  <h2>Experiment Configuration</h2>
  <table>
    <tr><th>Parameter</th><th>Value</th></tr>
    <tr><td>Number domain (n_max)</td><td>{cfg['n_max']}</td></tr>
    <tr><td>Draw size (k_count)</td><td>{cfg['k_count']}</td></tr>
    <tr><td>Order matters</td><td>{cfg['order_matters']}</td></tr>
    <tr><td>Description</td><td>{_fmt(cfg['description'])}</td></tr>
  </table>
  <h2>Summary</h2>
  <table>
    <tr><th>Key</th><th>Value</th></tr>
    <tr><td>Modules run</td><td>{_fmt(summary['modules_run'])}</td></tr>
    <tr><td>Models trained</td><td>{_fmt(summary['models_trained'])}</td></tr>
    <tr><td>Best model</td><td>{_fmt(summary['best_model'])}</td></tr>
    <tr><td>Best val score</td><td>{_fmt(summary['best_val_score'])}</td></tr>
    <tr><td>Ensemble prediction</td><td>{_fmt(summary['ensemble_prediction'])}</td></tr>
    <tr><td>Ensemble confidence</td><td>{_fmt(summary['ensemble_confidence'])}</td></tr>
  </table>
  <h2>Analysis Modules</h2>
  <table>
    <tr><th>Module</th><th>Status</th><th>Duration (s)</th></tr>
    {analysis_rows}
  </table>
  <h2>Trained Models</h2>
  <table>
    <tr><th>Model</th><th>Val Score</th><th>Test Score</th><th>Trained At</th></tr>
    {model_rows}
  </table>
</body>
</html>"""

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(html, encoding="utf-8")
    logger.info("HTML report written to %s", output_path)


def render_excel(payload: Dict[str, Any], output_path: Path) -> None:
    """
    Write the report payload to a multi-sheet Excel workbook.

    Requires ``openpyxl``.
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise ImportError(
            "openpyxl is required for Excel reports. Install with: pip install openpyxl"
        ) from exc

    wb = openpyxl.Workbook()
    header_fill = PatternFill(fill_type="solid", fgColor="1A56DB")
    header_font = Font(bold=True, color="FFFFFF")

    def _hdr(ws: Any) -> None:
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font

    # Summary sheet
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Parameter", "Value"])
    _hdr(ws)
    meta = payload["report_metadata"]
    cfg = payload["experiment_config"]
    summary = payload["summary"]
    for row in [
        ["Generated At", meta["generated_at"]],
        ["Experiment ID", meta["experiment_id"]],
        ["Experiment Name", meta["experiment_name"]],
        ["Status", meta["experiment_status"]],
        ["n_max", cfg["n_max"]],
        ["k_count", cfg["k_count"]],
        ["Order Matters", cfg["order_matters"]],
        ["Best Model", summary["best_model"]],
        ["Best Val Score", summary["best_val_score"]],
        ["Ensemble Prediction", str(summary["ensemble_prediction"])],
        ["Ensemble Confidence", summary["ensemble_confidence"]],
    ]:
        ws.append(row)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 50

    # Models sheet
    ws_m = wb.create_sheet("Models")
    ws_m.append(["Model Name", "Module", "Val Score", "Test Score", "Trained At"])
    _hdr(ws_m)
    for m in payload["models"]:
        ws_m.append([m["model_name"], m.get("module"), m.get("val_score"), m.get("test_score"), m.get("trained_at")])
    for col in ["A", "B", "C", "D", "E"]:
        ws_m.column_dimensions[col].width = 28

    # Analysis sheet
    ws_a = wb.create_sheet("Analysis")
    ws_a.append(["Module", "Status", "Duration (s)", "Created At"])
    _hdr(ws_a)
    for name, data in payload["analysis"].items():
        ws_a.append([name, data.get("status"), data.get("duration_seconds"), data.get("created_at")])
    for col in ["A", "B", "C", "D"]:
        ws_a.column_dimensions[col].width = 28

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))
    logger.info("Excel report written to %s", output_path)


def render_pdf(payload: Dict[str, Any], output_path: Path) -> Path:
    """
    Render the report as a PDF using WeasyPrint.

    Falls back to HTML if WeasyPrint is unavailable.

    Returns
    -------
    Path
        The actual output path (may differ from *output_path* on fallback).
    """
    html_path = output_path.with_suffix(".html")
    render_html(payload, html_path)
    try:
        from weasyprint import HTML as WeasyprintHTML
        WeasyprintHTML(filename=str(html_path)).write_pdf(str(output_path))
        logger.info("PDF report written to %s", output_path)
        return output_path
    except ImportError:
        logger.warning("weasyprint not installed; PDF report saved as HTML at %s", html_path)
        return html_path

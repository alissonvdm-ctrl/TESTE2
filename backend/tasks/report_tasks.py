"""
Report Celery tasks – routed to the ``reports`` queue.

Tasks
-----
- generate_report : Compile all persisted AnalysisResult / ModelResult rows
                    for an experiment into a structured document (JSON, HTML,
                    PDF, or Excel).

Time limits (from celery_config):
  soft 10 min / hard 20 min
"""

from __future__ import annotations

import json
import logging
import os
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

from celery.exceptions import SoftTimeLimitExceeded
from sqlalchemy import select
from sqlalchemy.orm import Session

from backend.core.config import settings
from backend.core.database import SyncSessionLocal
from backend.db.models import (
    AnalysisResult,
    Experiment,
    ModelResult,
    Report,
    ReportType,
)
from backend.tasks.celery_app import celery_app

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _load_experiment(session: Session, experiment_id: int) -> Experiment:
    exp = session.get(Experiment, experiment_id)
    if exp is None:
        raise ValueError(f"Experiment {experiment_id} not found")
    return exp


def _load_analysis_results(session: Session, experiment_id: int) -> List[AnalysisResult]:
    return list(
        session.scalars(
            select(AnalysisResult).where(
                AnalysisResult.experiment_id == experiment_id
            )
        ).all()
    )


def _load_model_results(session: Session, experiment_id: int) -> List[ModelResult]:
    return list(
        session.scalars(
            select(ModelResult).where(
                ModelResult.experiment_id == experiment_id
            )
        ).all()
    )


def _ensure_report_dir(experiment_id: int) -> Path:
    """Return (and create if missing) the report output directory."""
    report_dir = settings.upload_dir / "reports" / str(experiment_id)
    report_dir.mkdir(parents=True, exist_ok=True)
    return report_dir


def _build_report_payload(
    exp: Experiment,
    analysis_results: List[AnalysisResult],
    model_results: List[ModelResult],
) -> Dict[str, Any]:
    """
    Assemble the canonical report payload dict from DB records.

    This is the data structure written to JSON reports and used as the source
    for HTML / PDF rendering.
    """
    now_iso = datetime.now(tz=timezone.utc).isoformat()

    analysis_section: Dict[str, Any] = {}
    for ar in analysis_results:
        analysis_section[ar.module_name] = {
            "status": ar.status.value if ar.status else None,
            "duration_seconds": ar.duration_seconds,
            "created_at": ar.created_at.isoformat() if ar.created_at else None,
            "result": ar.result_json,
        }

    models_section: List[Dict[str, Any]] = []
    for mr in model_results:
        models_section.append(
            {
                "model_name": mr.model_name,
                "module": mr.module,
                "val_score": mr.val_score,
                "test_score": mr.test_score,
                "trained_at": mr.trained_at.isoformat() if mr.trained_at else None,
                "metrics": mr.metrics_json,
                "feature_importance": mr.feature_importance_json,
            }
        )

    # Best model summary
    scored_models = [m for m in models_section if m["val_score"] is not None]
    best_model = (
        min(scored_models, key=lambda m: m["val_score"])  # lower log-loss = better
        if scored_models
        else None
    )

    # Ensemble prediction (if available)
    ensemble_result = analysis_section.get("ensemble", {}).get("result") or {}

    return {
        "report_metadata": {
            "generated_at": now_iso,
            "experiment_id": exp.id,
            "experiment_name": exp.name,
            "experiment_status": exp.status.value if exp.status else None,
        },
        "experiment_config": {
            "n_max": exp.n_max,
            "k_count": exp.k_count,
            "order_matters": exp.order_matters,
            "description": exp.description,
            "config_json": exp.config_json,
        },
        "summary": {
            "modules_run": list(analysis_section.keys()),
            "models_trained": [m["model_name"] for m in models_section],
            "best_model": best_model["model_name"] if best_model else None,
            "best_val_score": best_model["val_score"] if best_model else None,
            "ensemble_prediction": ensemble_result.get("predicted_numbers"),
            "ensemble_confidence": ensemble_result.get("confidence"),
        },
        "analysis": analysis_section,
        "models": models_section,
    }


def _write_json_report(payload: Dict[str, Any], output_path: Path) -> None:
    """Write the report payload to a JSON file."""
    with open(output_path, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, indent=2, ensure_ascii=False, default=str)


def _write_html_report(payload: Dict[str, Any], output_path: Path) -> None:
    """Render a minimal but complete HTML report from the payload dict."""
    meta = payload["report_metadata"]
    cfg = payload["experiment_config"]
    summary = payload["summary"]
    analysis = payload["analysis"]
    models = payload["models"]

    def _fmt(v: Any) -> str:
        if v is None:
            return "<em>N/A</em>"
        if isinstance(v, float):
            return f"{v:.4f}"
        if isinstance(v, list):
            return ", ".join(str(x) for x in v)
        return str(v)

    # Build model rows
    model_rows = "".join(
        f"<tr><td>{m['model_name']}</td>"
        f"<td>{_fmt(m['val_score'])}</td>"
        f"<td>{_fmt(m['test_score'])}</td>"
        f"<td>{_fmt(m['trained_at'])}</td></tr>"
        for m in models
    )

    # Build analysis rows
    analysis_rows = "".join(
        f"<tr><td>{name}</td>"
        f"<td>{_fmt(data.get('status'))}</td>"
        f"<td>{_fmt(data.get('duration_seconds'))}</td></tr>"
        for name, data in analysis.items()
    )

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8"/>
  <title>NumericSequenceAnalyzer – Report #{meta['experiment_id']}</title>
  <style>
    body {{ font-family: sans-serif; margin: 2rem; color: #222; }}
    h1 {{ color: #1a56db; }}
    h2 {{ border-bottom: 2px solid #e5e7eb; padding-bottom: 0.3rem; }}
    table {{ border-collapse: collapse; width: 100%; margin-bottom: 1.5rem; }}
    th, td {{ border: 1px solid #d1d5db; padding: 0.5rem 0.75rem; text-align: left; }}
    th {{ background: #f3f4f6; }}
    .badge {{ display:inline-block; padding:0.2rem 0.6rem; border-radius:9999px;
               background:#dbeafe; color:#1e40af; font-size:0.85rem; }}
  </style>
</head>
<body>
  <h1>Analysis Report – {meta['experiment_name']}</h1>
  <p>Generated: {meta['generated_at']} &nbsp;|&nbsp; Status:
     <span class="badge">{meta['experiment_status'] or 'unknown'}</span></p>

  <h2>Experiment Configuration</h2>
  <table>
    <tr><th>Parameter</th><th>Value</th></tr>
    <tr><td>Number domain (n_max)</td><td>{cfg['n_max']}</td></tr>
    <tr><td>Draw size (k_count)</td><td>{cfg['k_count']}</td></tr>
    <tr><td>Order matters</td><td>{cfg['order_matters']}</td></tr>
    <tr><td>Description</td><td>{_fmt(cfg['description'])}</td></tr>
  </table>

  <h2>Summary</h2>
  <table>
    <tr><th>Key</th><th>Value</th></tr>
    <tr><td>Modules run</td><td>{_fmt(summary['modules_run'])}</td></tr>
    <tr><td>Models trained</td><td>{_fmt(summary['models_trained'])}</td></tr>
    <tr><td>Best model</td><td>{_fmt(summary['best_model'])}</td></tr>
    <tr><td>Best val score</td><td>{_fmt(summary['best_val_score'])}</td></tr>
    <tr><td>Ensemble prediction</td><td>{_fmt(summary['ensemble_prediction'])}</td></tr>
    <tr><td>Ensemble confidence</td><td>{_fmt(summary['ensemble_confidence'])}</td></tr>
  </table>

  <h2>Analysis Modules</h2>
  <table>
    <tr><th>Module</th><th>Status</th><th>Duration (s)</th></tr>
    {analysis_rows}
  </table>

  <h2>Trained Models</h2>
  <table>
    <tr><th>Model</th><th>Val Score</th><th>Test Score</th><th>Trained At</th></tr>
    {model_rows}
  </table>
</body>
</html>"""

    output_path.write_text(html, encoding="utf-8")


# ---------------------------------------------------------------------------
# Task: generate_report
# ---------------------------------------------------------------------------


@celery_app.task(
    name="tasks.generate_report",
    bind=True,
    max_retries=2,
    default_retry_delay=30,
    soft_time_limit=600,   # 10 min
    time_limit=1200,       # 20 min
    acks_late=True,
)
def generate_report(
    self,
    experiment_id: int,
    report_format: str = "json",
) -> Dict[str, Any]:
    """
    Compile all results for an experiment into a report file.

    Supported formats
    -----------------
    - ``json``  – machine-readable JSON (always available)
    - ``html``  – self-contained HTML page (always available)
    - ``pdf``   – PDF rendered from the HTML page (requires ``weasyprint``)
    - ``excel`` – Excel workbook (requires ``openpyxl``)

    The generated file is written to
    ``{settings.upload_dir}/reports/{experiment_id}/{timestamp}.{ext}``
    and a ``Report`` row is inserted into the database with the file path.

    Parameters
    ----------
    experiment_id : int
        Primary key of the Experiment.
    report_format : str
        One of ``'json'``, ``'html'``, ``'pdf'``, ``'excel'``.

    Returns
    -------
    dict
        ``{experiment_id, report_format, file_path, duration_seconds}``
    """
    t0 = time.perf_counter()
    fmt = report_format.lower()
    logger.info("generate_report START: experiment_id=%s format=%s", experiment_id, fmt)

    _valid_formats = {"json", "html", "pdf", "excel"}
    if fmt not in _valid_formats:
        raise ValueError(f"Unsupported report format {fmt!r}. Choose from {_valid_formats}.")

    try:
        with SyncSessionLocal() as session:
            exp = _load_experiment(session, experiment_id)
            analysis_results = _load_analysis_results(session, experiment_id)
            model_results = _load_model_results(session, experiment_id)

            payload = _build_report_payload(exp, analysis_results, model_results)

            report_dir = _ensure_report_dir(experiment_id)
            timestamp = datetime.now(tz=timezone.utc).strftime("%Y%m%dT%H%M%S")

            if fmt == "json":
                file_path = report_dir / f"{timestamp}.json"
                _write_json_report(payload, file_path)
                report_type = ReportType.JSON

            elif fmt == "html":
                file_path = report_dir / f"{timestamp}.html"
                _write_html_report(payload, file_path)
                report_type = ReportType.HTML

            elif fmt == "pdf":
                # Render HTML first, then convert to PDF with weasyprint
                html_path = report_dir / f"{timestamp}.html"
                _write_html_report(payload, html_path)
                file_path = report_dir / f"{timestamp}.pdf"
                try:
                    from weasyprint import HTML as WeasyprintHTML
                    WeasyprintHTML(filename=str(html_path)).write_pdf(str(file_path))
                except ImportError:
                    logger.warning(
                        "weasyprint not installed; falling back to HTML for experiment_id=%s",
                        experiment_id,
                    )
                    file_path = html_path
                    fmt = "html"
                report_type = ReportType.PDF if fmt == "pdf" else ReportType.HTML

            elif fmt == "excel":
                file_path = report_dir / f"{timestamp}.xlsx"
                _write_excel_report(payload, file_path)
                report_type = ReportType.EXCEL

            # Persist Report row
            report_record = Report(
                experiment_id=experiment_id,
                report_type=report_type,
                file_path=str(file_path),
            )
            session.add(report_record)
            session.commit()
            session.refresh(report_record)
            report_id = report_record.id

        duration = time.perf_counter() - t0

        logger.info(
            "generate_report DONE: experiment_id=%s format=%s file=%s in %.2fs",
            experiment_id, fmt, file_path, duration,
        )
        return {
            "experiment_id": experiment_id,
            "report_id": report_id,
            "report_format": fmt,
            "file_path": str(file_path),
            "duration_seconds": duration,
        }

    except SoftTimeLimitExceeded:
        logger.warning("generate_report soft time limit exceeded: experiment_id=%s", experiment_id)
        raise

    except Exception as exc:
        logger.exception("generate_report FAILED: experiment_id=%s", experiment_id)
        raise self.retry(exc=exc)


# ---------------------------------------------------------------------------
# Excel helper
# ---------------------------------------------------------------------------


def _write_excel_report(payload: Dict[str, Any], output_path: Path) -> None:
    """
    Write the report payload to a multi-sheet Excel workbook.

    Sheets
    ------
    - Summary     – high-level experiment metadata and results
    - Models      – one row per trained model with scores
    - Analysis    – one row per analysis module with status / duration

    Requires ``openpyxl``.  Raises ``ImportError`` if not installed.
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:  # pragma: no cover
        raise ImportError(
            "openpyxl is required to generate Excel reports. "
            "Install it with: pip install openpyxl"
        ) from exc

    wb = openpyxl.Workbook()

    # ---- Summary sheet -------------------------------------------------------
    ws_summary = wb.active
    ws_summary.title = "Summary"
    header_fill = PatternFill(fill_type="solid", fgColor="1A56DB")
    header_font = Font(bold=True, color="FFFFFF")

    summary_rows = [
        ["Generated At", payload["report_metadata"]["generated_at"]],
        ["Experiment ID", payload["report_metadata"]["experiment_id"]],
        ["Experiment Name", payload["report_metadata"]["experiment_name"]],
        ["Status", payload["report_metadata"]["experiment_status"]],
        ["n_max", payload["experiment_config"]["n_max"]],
        ["k_count", payload["experiment_config"]["k_count"]],
        ["Order Matters", payload["experiment_config"]["order_matters"]],
        ["Best Model", payload["summary"]["best_model"]],
        ["Best Val Score", payload["summary"]["best_val_score"]],
        ["Ensemble Prediction", str(payload["summary"]["ensemble_prediction"])],
        ["Ensemble Confidence", payload["summary"]["ensemble_confidence"]],
    ]
    ws_summary.append(["Parameter", "Value"])
    for cell in ws_summary[1]:
        cell.fill = header_fill
        cell.font = header_font
    for row in summary_rows:
        ws_summary.append(row)
    ws_summary.column_dimensions["A"].width = 30
    ws_summary.column_dimensions["B"].width = 50

    # ---- Models sheet -------------------------------------------------------
    ws_models = wb.create_sheet("Models")
    ws_models.append(["Model Name", "Module", "Val Score", "Test Score", "Trained At"])
    for cell in ws_models[1]:
        cell.fill = header_fill
        cell.font = header_font
    for m in payload["models"]:
        ws_models.append([
            m["model_name"],
            m["module"],
            m["val_score"],
            m["test_score"],
            m["trained_at"],
        ])
    for col in ["A", "B", "C", "D", "E"]:
        ws_models.column_dimensions[col].width = 30

    # ---- Analysis sheet -----------------------------------------------------
    ws_analysis = wb.create_sheet("Analysis")
    ws_analysis.append(["Module", "Status", "Duration (s)", "Created At"])
    for cell in ws_analysis[1]:
        cell.fill = header_fill
        cell.font = header_font
    for name, data in payload["analysis"].items():
        ws_analysis.append([
            name,
            data.get("status"),
            data.get("duration_seconds"),
            data.get("created_at"),
        ])
    for col in ["A", "B", "C", "D"]:
        ws_analysis.column_dimensions[col].width = 30

    wb.save(str(output_path))
